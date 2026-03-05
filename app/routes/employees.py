"""
Employees list and management
"""
import io
import os
from datetime import datetime, timedelta
from pathlib import Path

from flask import Blueprint, render_template, request, redirect, url_for, flash, abort, current_app, send_file
from flask_login import login_required, current_user
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.barcharts import VerticalBarChart

from app import db
from app.models import User, Department, Recognition
from app.models.goal import Goal
from app.models.review import Review
from app.models.user import Role
from app.utils import admin_required, manager_required
from app.utils.performance import compute_user_performance_series

employees_bp = Blueprint("employees", __name__)


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in current_app.config.get("ALLOWED_EXTENSIONS", set())


@employees_bp.route("/")
@login_required
def index():
    q = request.args.get("q", "").strip()
    dept_id = request.args.get("department", type=int)
    page = request.args.get("page", 1, type=int)
    # By default show only active; admin can optionally include inactive
    include_inactive = current_user.is_admin() and request.args.get("inactive") == "1"
    # Everyone sees the org directory; visibility of details/actions is controlled elsewhere.
    query = User.query
    if not include_inactive:
        query = query.filter_by(is_active=True)
    # Hide admin accounts from non-admin viewers
    if not current_user.is_admin():
        query = query.filter(User.role != Role.ADMIN)
    # Employees and managers see only their own department; admin sees all
    if not current_user.is_admin() and current_user.department_id is not None:
        query = query.filter(User.department_id == current_user.department_id)
        dept_id = current_user.department_id  # ignore URL dept filter for non-admin
    if q:
        query = query.filter(
            db.or_(
                User.first_name.ilike(f"%{q}%"),
                User.last_name.ilike(f"%{q}%"),
                User.email.ilike(f"%{q}%"),
            )
        )
    if dept_id:
        query = query.filter(User.department_id == dept_id)
    query = query.order_by(User.first_name, User.last_name)
    pagination = query.paginate(page=page, per_page=12)
    # Non-admin only see their own department in the filter dropdown
    if current_user.is_admin():
        departments = Department.query.order_by(Department.name).all()
    else:
        departments = [current_user.department] if current_user.department else []
    return render_template(
        "employees/index.html",
        pagination=pagination,
        departments=departments,
        q=q,
        dept_id=dept_id,
        include_inactive=include_inactive,
    )


@employees_bp.route("/<int:user_id>")
@login_required
def view(user_id):
    user = User.query.get_or_404(user_id)
    if not current_user.can_manage_user(user) and current_user.id != user_id:
        abort(403)
    recognitions = user.recognitions_received.order_by(Recognition.created_at.desc()).limit(10).all()
    return render_template("employees/view.html", user=user, recognitions=recognitions)


@employees_bp.route("/<int:user_id>/export/excel")
@login_required
@manager_required
def export_employee_performance_excel(user_id):
    """Export an employee's goals and reviews to an Excel workbook."""
    user = User.query.get_or_404(user_id)
    if not current_user.can_manage_user(user) and current_user.id != user_id:
        abort(403)

    goals = Goal.query.filter_by(user_id=user.id).order_by(Goal.target_date, Goal.created_at).all()
    reviews = Review.query.filter_by(user_id=user.id).order_by(Review.period_start, Review.created_at).all()

    wb = Workbook()
    ws_goals = wb.active
    ws_goals.title = "Goals"

    ws_goals.append(
        [
            "Title",
            "Description",
            "Target Date",
            "Status",
            "Progress (%)",
            "Weight",
            "Created At",
            "Updated At",
        ]
    )
    for g in goals:
        ws_goals.append(
            [
                g.title,
                g.description or "",
                g.target_date.isoformat() if g.target_date else "",
                g.status,
                g.progress,
                g.weight,
                g.created_at.isoformat() if g.created_at else "",
                g.updated_at.isoformat() if g.updated_at else "",
            ]
        )

    ws_reviews = wb.create_sheet("Reviews")
    ws_reviews.append(
        [
            "Period Start",
            "Period End",
            "Reviewer",
            "Rating",
            "Status",
            "Submitted At",
            "Summary",
        ]
    )
    for r in reviews:
        ws_reviews.append(
            [
                r.period_start.isoformat() if r.period_start else "",
                r.period_end.isoformat() if r.period_end else "",
                r.reviewer.full_name if r.reviewer else "",
                r.rating if r.rating is not None else "",
                r.status,
                r.submitted_at.isoformat() if r.submitted_at else "",
                (r.summary or "")[:500],
            ]
        )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"{user.username}_performance.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@employees_bp.route("/<int:user_id>/export/pdf")
@login_required
@manager_required
def export_employee_performance_pdf(user_id):
    """Export an employee's goals, reviews, and ratings to a PDF report."""
    user = User.query.get_or_404(user_id)
    if not current_user.can_manage_user(user) and current_user.id != user_id:
        abort(403)

    goals = Goal.query.filter_by(user_id=user.id).order_by(Goal.target_date, Goal.created_at).all()
    reviews = Review.query.filter_by(user_id=user.id).order_by(Review.period_start, Review.created_at).all()

    buffer = io.BytesIO()
    page_w, page_h = A4
    report_id = f"MF-{user.id}-{datetime.utcnow().strftime('%Y%m%d%H%M')}"
    generated_ts = datetime.utcnow().strftime("%Y-%m-%d %H:%M") + " UTC"

    def add_pdf_header_footer(canvas, doc):
        canvas.saveState()
        # Left accent bar (full height) – unique branding
        canvas.setFillColor(colors.HexColor("#0ea5e9"))
        canvas.rect(0, 0, 4 * mm, page_h, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor("#38bdf8"))
        canvas.rect(0, 0, 2 * mm, page_h, fill=1, stroke=0)
        # Header: two-tone strip with report ID
        canvas.setFillColor(colors.HexColor("#0f172a"))
        canvas.rect(0, page_h - 24 * mm, page_w, 24 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor("#0ea5e9"))
        canvas.rect(0, page_h - 24 * mm, page_w * 0.45, 24 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 14)
        canvas.drawString(20 * mm, page_h - 11 * mm, "MetricForge")
        canvas.setFont("Helvetica", 9)
        canvas.drawString(20 * mm, page_h - 15 * mm, "Employee Performance Report · Confidential")
        canvas.drawString(20 * mm, page_h - 18 * mm, "Generated " + generated_ts)
        canvas.drawString(20 * mm, page_h - 21 * mm, "Report ID: " + report_id)
        # Watermark (subtle, behind content)
        canvas.setFillColor(colors.HexColor("#e5e7eb"))
        canvas.setFillAlpha(0.12)
        canvas.setFont("Helvetica-Bold", 72)
        canvas.rotate(45)
        canvas.drawCentredString(page_w / 2 + 40 * mm, -page_h / 2 + 40 * mm, "CONFIDENTIAL")
        canvas.rotate(-45)
        canvas.setFillAlpha(1)
        # Footer: accent line + page number + report ID
        canvas.setFillColor(colors.HexColor("#0ea5e9"))
        canvas.rect(20 * mm, 18 * mm, page_w - 40 * mm, 2 * mm, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(page_w - 20 * mm, 12 * mm, f"Page {doc.page} · " + user.full_name)
        canvas.drawString(20 * mm, 12 * mm, "MetricForge EPMS · " + report_id)
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=26 * mm,
        bottomMargin=26 * mm,
        title=f"{user.full_name} – Performance Report",
        onFirstPage=add_pdf_header_footer,
        onLaterPages=add_pdf_header_footer,
    )
    styles = getSampleStyleSheet()
    title_style = styles["Heading1"]
    title_style.fontName = "Helvetica-Bold"
    title_style.fontSize = 18
    title_style.textColor = colors.HexColor("#0ea5e9")

    subtitle_style = styles["Heading3"]
    subtitle_style.textColor = colors.HexColor("#4b5563")

    normal = styles["Normal"]
    normal.fontSize = 9

    small = styles["Normal"]
    small.fontSize = 8
    small.textColor = colors.HexColor("#6b7280")

    story = []

    # Cover / title section with report ID and confidentiality
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("EMPLOYEE PERFORMANCE REPORT", subtitle_style))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph(f"<b>{user.full_name}</b>", title_style))
    story.append(Paragraph(f"<i>{user.job_title or 'Employee'}</i> · {user.department.name if user.department else '—'}", normal))
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph(f"<b>Report ID:</b> {report_id}", small))
    story.append(Paragraph(f"<b>Generated:</b> {generated_ts}", small))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("This document is confidential and intended for internal HR and management use only. Do not distribute without authorization.", small))
    story.append(Spacer(1, 12 * mm))

    # Report title (section)
    story.append(Paragraph("<b>1. Employee details</b>", subtitle_style))
    story.append(Spacer(1, 4 * mm))
    meta_lines = [
        f"Employee: {user.full_name}",
        f"Email: {user.email}",
        f"Department: {user.department.name if user.department else '—'}",
        f"Job title: {user.job_title or 'Employee'}",
        f"Username: {user.username}",
    ]
    for line in meta_lines:
        story.append(Paragraph(line, normal))

    # Key metrics box – eye-catching summary
    completed_goals = sum(1 for g in goals if g.status == "completed")
    ratings_with_val = [r.rating for r in reviews if r.rating is not None]
    avg_rating = (sum(ratings_with_val) / len(ratings_with_val)) if ratings_with_val else None
    metrics_data = [
        ["Goals", "Completed", "Reviews", "Avg rating"],
        [str(len(goals)), str(completed_goals), str(len(reviews)), f"{avg_rating:.1f}/5" if avg_rating is not None else "—"],
    ]
    metrics_table = Table(metrics_data, colWidths=[35 * mm, 35 * mm, 35 * mm, 35 * mm])
    metrics_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0ea5e9")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#f0f9ff")),
            ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#0c4a6e")),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#0ea5e9")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#bae6fd")),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ])
    )
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("<b>Key metrics at a glance</b>", subtitle_style))
    story.append(Spacer(1, 2 * mm))
    story.append(metrics_table)
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("All information in this report is confidential and intended for internal HR use.", small))
    story.append(Spacer(1, 10 * mm))

    # Goals table
    story.append(Paragraph("<b>2. Goals overview</b>", subtitle_style))
    story.append(Spacer(1, 2 * mm))

    if not goals:
        story.append(Paragraph("No goals recorded for this employee.", normal))
    else:
        goals_data = [
            [
                "Title",
                "Status",
                "Progress",
                "Target date",
            ]
        ]
        for g in goals:
            target_date = g.target_date.isoformat() if g.target_date else "—"
            goals_data.append(
                [
                    g.title,
                    g.status.title(),
                    f"{g.progress}%",
                    target_date,
                ]
            )

        goals_table = Table(goals_data, colWidths=[60 * mm, 25 * mm, 25 * mm, 35 * mm])
        goals_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 9),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.whitesmoke),
                    ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#111827")),
                    ("ALIGN", (1, 1), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1d5db")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.HexColor("#e5e7eb")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(goals_table)

    story.append(Spacer(1, 10 * mm))

    # Reviews table
    story.append(Paragraph("<b>3. Performance reviews &amp; ratings</b>", subtitle_style))
    story.append(Spacer(1, 2 * mm))

    if not reviews:
        story.append(Paragraph("No performance reviews have been recorded.", normal))
    else:
        reviews_data = [
            [
                "Period",
                "Reviewer",
                "Rating",
                "Status",
            ]
        ]
        for r in reviews:
            if r.period_start and r.period_end:
                period = f"{r.period_start.isoformat()} → {r.period_end.isoformat()}"
            else:
                period = "—"
            reviewer_name = r.reviewer.full_name if r.reviewer else "N/A"
            rating_text = f"{r.rating}/5" if r.rating is not None else "Not rated"
            reviews_data.append(
                [
                    period,
                    reviewer_name,
                    rating_text,
                    r.status.title(),
                ]
            )

        reviews_table = Table(reviews_data, colWidths=[45 * mm, 55 * mm, 25 * mm, 25 * mm])
        reviews_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 9),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.whitesmoke),
                    ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#111827")),
                    ("ALIGN", (0, 1), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1d5db")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.HexColor("#e5e7eb")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(reviews_table)

    # Composite performance trend chart (same formula as dashboard)
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("<b>4. Overall performance trend</b> (last 12 months)", subtitle_style))
    story.append(Spacer(1, 2 * mm))

    trend_start = datetime.utcnow() - timedelta(days=365)
    perf_series = compute_user_performance_series(
        user_id=user.id,
        start_date=trend_start,
        end_date=None,
        max_points=12,
    )

    if not perf_series:
        story.append(Paragraph("Not enough data to compute a performance trend for this employee.", normal))
    else:
        labels = [label for label, _ in perf_series]
        values = [score for _, score in perf_series]

        chart_width = 160 * mm
        chart_height = 70 * mm
        drawing = Drawing(chart_width, chart_height)
        chart = VerticalBarChart()
        chart.x = 20
        chart.y = 20
        chart.height = chart_height - 40
        chart.width = chart_width - 40
        chart.data = [values]
        chart.categoryAxis.categoryNames = labels
        chart.valueAxis.valueMin = 0
        chart.valueAxis.valueMax = 100
        chart.valueAxis.valueStep = 20
        chart.barWidth = 10
        chart.bars[0].fillColor = colors.HexColor("#0ea5e9")
        chart.bars[0].strokeColor = colors.HexColor("#0284c7")
        chart.bars[0].strokeWidth = 0.5
        chart.categoryAxis.labels.fontName = "Helvetica"
        chart.valueAxis.labels.fontName = "Helvetica"
        drawing.add(chart)
        story.append(drawing)

    story.append(Spacer(1, 8 * mm))
    story.append(
        Paragraph(
            "<i>Prepared by MetricForge – Employee Performance Management System. This document is confidential. Report ID: " + report_id + "</i>",
            small,
        )
    )

    doc.build(story)

    buffer.seek(0)
    filename = f"{user.username}_performance.pdf"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf",
    )


@employees_bp.route("/<int:user_id>/edit", methods=["GET", "POST"])
@login_required
def edit(user_id):
    user = User.query.get_or_404(user_id)
    if not current_user.can_manage_user(user):
        abort(403)
    if request.method == "POST":
        user.first_name = request.form.get("first_name", "").strip()
        user.last_name = request.form.get("last_name", "").strip()
        user.job_title = request.form.get("job_title", "").strip() or None
        user.phone = request.form.get("phone", "").strip() or None
        user.department_id = request.form.get("department_id", type=int) or None
        manager_id = request.form.get("manager_id", type=int)
        if manager_id is not None:
            user.manager_id = manager_id if manager_id else None
        if current_user.is_admin():
            role = request.form.get("role", "").strip()
            if role in (Role.ADMIN, Role.MANAGER, Role.EMPLOYEE):
                user.role = role
        # Avatar upload
        if "avatar" in request.files:
            f = request.files["avatar"]
            if f and f.filename and allowed_file(f.filename):
                upload_dir = Path(current_app.config["UPLOAD_FOLDER"]) / "avatars"
                upload_dir.mkdir(parents=True, exist_ok=True)
                ext = f.filename.rsplit(".", 1)[1].lower()
                filename = f"user_{user.id}_{os.urandom(4).hex()}.{ext}"
                filepath = upload_dir / filename
                f.save(str(filepath))
                user.avatar_url = f"uploads/avatars/{filename}"
        db.session.commit()
        flash("Member updated.", "success")
        return redirect(url_for("employees.view", user_id=user.id))
    departments = Department.query.order_by(Department.name).all()
    managers = User.query.filter(User.id != user.id, User.is_active).order_by(User.first_name, User.last_name).all()
    return render_template("employees/edit.html", user=user, departments=departments, managers=managers)


@employees_bp.route("/<int:user_id>/delete", methods=["POST"])
@login_required
def delete(user_id):
    user = User.query.get_or_404(user_id)
    if not current_user.can_manage_user(user):
        abort(403)
    if user.id == current_user.id:
        flash("You cannot deactivate your own account.", "error")
        return redirect(url_for("employees.view", user_id=user.id))
    user.is_active = False
    db.session.commit()
    flash("Member deactivated. They can no longer sign in.", "info")
    return redirect(url_for("employees.index"))


@employees_bp.route("/create-manager", methods=["GET", "POST"])
@login_required
@admin_required
def create_manager():
    """Admin-only: create a new manager account."""
    from werkzeug.security import generate_password_hash

    departments = Department.query.order_by(Department.name).all()
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        username = request.form.get("username", "").strip()
        first_name = request.form.get("first_name", "").strip()
        last_name = request.form.get("last_name", "").strip()
        job_title = request.form.get("job_title", "").strip() or "Manager"
        department_id = request.form.get("department_id", type=int) or None
        password = request.form.get("password", "")

        if not email or not username or not first_name or not last_name:
            flash("All fields are required.", "error")
        elif User.query.filter_by(email=email).first():
            flash("Email already registered.", "error")
        elif User.query.filter_by(username=username).first():
            flash("Username already taken.", "error")
        elif len(password) < 8:
            flash("Password must be at least 8 characters.", "error")
        else:
            user = User(
                email=email,
                username=username,
                first_name=first_name,
                last_name=last_name,
                role=Role.MANAGER,
                department_id=department_id,
                manager_id=current_user.id,
                job_title=job_title,
            )
            user.set_password(password)
            db.session.add(user)
            db.session.commit()
            flash("Manager account created.", "success")
            return redirect(url_for("employees.view", user_id=user.id))

    return render_template("employees/create_manager.html", departments=departments)
